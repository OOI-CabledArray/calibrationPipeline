"""List the CTD casts of each cruise made at each OOI site, from the CastList sheet of the cast-log workbook.

  python scripts/site_casts.py                    every cruise with a cast-log workbook named in the config
  python scripts/site_casts.py --site RS01SBPS    casts at one site

Prints cast, site, cast flag and the number of salinity samples listed for the cast.
"""
from __future__ import annotations

import argparse
import sys
import warnings
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from rca_sal import github_data  # noqa: E402
from rca_sal.pipeline import load_config  # noqa: E402


def cast_list(path) -> list[dict]:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        ws = openpyxl.load_workbook(path, data_only=True, read_only=True)["CastList"]
        rows = list(ws.iter_rows(values_only=True))
    h = [str(c).strip() if c else "" for c in rows[0]]
    return [dict(zip(h, r)) for r in rows[1:] if r[h.index("Cast")]]


def salinity_sample_counts(path) -> dict[str, int]:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        rows = list(openpyxl.load_workbook(path, data_only=True, read_only=True)["SampleList"].iter_rows(values_only=True))
    h = [str(c).strip() if c else "" for c in rows[0]]
    ci, si = h.index("Cast"), h.index("Salinity Bottle Number")
    out: dict[str, int] = {}
    for r in rows[1:]:
        if r[ci] and r[si]:
            out[str(r[ci])] = out.get(str(r[ci]), 0) + 1
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--site")
    ap.add_argument("--config", default=str(ROOT / "config" / "cruises.toml"))
    a = ap.parse_args(argv)
    cfg = load_config(a.config)
    srcs = github_data.sources_from_config(cfg, ROOT)
    for year, y in sorted(cfg["years"].items()):
        spec = y.get("inputs", {}).get("cast_log_xlsm")
        if not spec:
            continue
        res, _ = github_data.resolve_inputs(srcs, {"cast_log_xlsm": spec})
        counts = salinity_sample_counts(res["cast_log_xlsm"])
        print(f"{year}  {y.get('cruise')} {y.get('ship_cruise_id')}")
        for r in cast_list(res["cast_log_xlsm"]):
            site = r.get("Target Asset")
            if str(r["Cast"]).startswith("CTD") and (not a.site or site == a.site):
                print(f"   {r['Cast']:8s} {str(site):12s} {str(r.get('Station')):36s} salinity samples {counts.get(str(r['Cast']), 0)}  [{r.get('Cast Flag')}]")
    return 0


if __name__ == "__main__":
    sys.exit(main())
